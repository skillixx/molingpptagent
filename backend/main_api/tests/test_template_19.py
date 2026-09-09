"""水彩绿植轻商务模板的探针、素材、语义和图片协议测试。"""

from __future__ import annotations

import json
import os
import re
import subprocess
import sys
from html import unescape
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation as PowerPointPresentation

from backend.main_api.workers.template_renderer import PresentationTemplateRenderer, TemplateRenderError


REPOSITORY_ROOT = Path(__file__).resolve().parents[3]
TEMPLATE_ROOT = Path(__file__).resolve().parents[1] / "template"
TEMPLATE_PATH = TEMPLATE_ROOT / "template_19.json"
BUILDER_PATH = REPOSITORY_ROOT / "utils" / "build_watercolor_botanical_template.mjs"
ROUNDTRIP_VERIFIER_PATH = REPOSITORY_ROOT / "utils" / "verify_template_19_pptx_roundtrip.mjs"

PROBE_IDS = {"cover-botanical-frame", "content-text-4", "content-image-1"}
MVP_IDS = {
    "cover-botanical-frame",
    "contents-2",
    "contents-3",
    "contents-4",
    "contents-5",
    "contents-6",
    "contents-10",
    "transition-eucalyptus",
    "content-text-2",
    "content-text-3",
    "content-text-4",
    "end-botanical-frame",
}
PRODUCTION_IDS = MVP_IDS | {
    "cover-image",
    "transition-fern",
    "content-statement-1",
    "content-image-1",
    "content-metrics-4",
    "end-action",
}
ASSET_CONTRACT = {
    "template_19_asset_bg_cover_v1.jpg": ((1920, 1080), "RGB", 450_000),
    "template_19_asset_bg_content_v1.jpg": ((1920, 1080), "RGB", 250_000),
    "template_19_asset_bg_section_v1.jpg": ((1920, 1080), "RGB", 350_000),
    "template_19_asset_bg_end_v1.jpg": ((1920, 1080), "RGB", 400_000),
    "template_19_asset_eucalyptus_corner_upper_v1.png": ((1400, 900), "RGBA", 800_000),
    "template_19_asset_eucalyptus_corner_v1.png": ((1400, 900), "RGBA", 800_000),
    "template_19_asset_eucalyptus_sweep_v1.png": ((1800, 650), "RGBA", 950_000),
    "template_19_asset_fern_spray_v1.png": ((1400, 950), "RGBA", 850_000),
    "template_19_asset_leaf_medallion_v1.png": ((900, 900), "RGBA", 600_000),
}


def _renderer() -> PresentationTemplateRenderer:
    return PresentationTemplateRenderer(TEMPLATE_ROOT)


def _template() -> dict:
    return json.loads(TEMPLATE_PATH.read_text(encoding="utf-8"))


def _slot_type(element: dict) -> str | None:
    value = element.get("textType")
    if isinstance(value, str):
        return value
    text = element.get("text")
    if isinstance(text, dict) and isinstance(text.get("type"), str):
        return text["type"]
    return None


def _plain_text(element: dict) -> str:
    value = element.get("content")
    if not isinstance(value, str):
        text = element.get("text")
        value = text.get("content") if isinstance(text, dict) else ""
    return unescape(re.sub(r"<[^>]+>", "", value or ""))


def _items(count: int) -> list[dict[str, str]]:
    return [
        {"title": f"完整业务标题第{index}项需要安全展示", "text": f"第{index}项完整说明。"}
        for index in range(1, count + 1)
    ]


def _images(count: int, *, width: int = 1600, height: int = 900) -> list[dict[str, object]]:
    return [
        {"src": f"https://example.invalid/template-19-{index}.jpg", "width": width, "height": height}
        for index in range(1, count + 1)
    ]


def test_template_19_probe_builder_is_deterministic(tmp_path: Path) -> None:
    """技术探针构建两次必须得到完全一致的三页JSON。"""

    outputs = [tmp_path / f"probe-{index}.json" for index in (1, 2)]
    for output in outputs:
        result = subprocess.run(
            ["node", str(BUILDER_PATH), "--stage", "probe", str(output)],
            cwd=REPOSITORY_ROOT,
            capture_output=True,
            text=True,
            timeout=30,
            check=False,
        )
        assert result.returncode == 0, result.stderr
    first = json.loads(outputs[0].read_text(encoding="utf-8"))
    second = json.loads(outputs[1].read_text(encoding="utf-8"))
    assert first == second
    assert first["metadata"]["buildStage"] == "probe"
    assert {slide["id"] for slide in first["slides"]} == PROBE_IDS


def test_template_19_probe_assets_are_valid() -> None:
    """探针实际使用的背景和桉叶角景必须满足发布格式。"""

    for filename in (
        "template_19_asset_bg_cover_v1.jpg",
        "template_19_asset_bg_content_v1.jpg",
        "template_19_asset_eucalyptus_corner_v1.png",
    ):
        size, mode, limit = ASSET_CONTRACT[filename]
        path = TEMPLATE_ROOT / filename
        with Image.open(path) as image:
            assert image.size == size
            assert image.mode == mode
            if mode == "RGBA":
                assert image.getchannel("A").getextrema() == (0, 255)
        assert path.stat().st_size <= limit


def test_template_19_probe_preserves_reference_contract() -> None:
    """探针必须保留参考页映射和水彩绿植视觉契约。"""

    template = json.loads(TEMPLATE_PATH.read_text(encoding="utf-8"))
    assert template["id"] == "template_19"
    assert template["title"] == "水彩绿植轻商务"
    assert (template["width"], template["height"]) == (1000, 562.5)
    assert set(template["metadata"]["probeSlideIds"]) == PROBE_IDS
    assert template["metadata"]["rightsPolicy"] == "reference-media-excluded"
    assert all(slide.get("sourceReferenceSlides") for slide in template["slides"])
    assert all(isinstance(slide.get("sourceFidelity"), dict) for slide in template["slides"])


def test_template_19_four_item_probe_keeps_four_item_density() -> None:
    """长原标题不能把四项探针静默退化为单项页。"""

    page = _renderer().render(
        template_id="template_19",
        semantic_slides=[{"type": "content", "data": {"title": "四项内容", "items": _items(4)}}],
        task_id="template-19-probe-four",
        fallback_title="四项内容",
    )["slides"][0]
    assert page["templateSlideId"] == "content-text-4"
    assert sum(element.get("textType") == "itemTitle" for element in page["elements"]) == 4


@pytest.mark.parametrize(("width", "height"), [(1600, 900), (900, 1600), (1000, 1000)])
def test_template_19_probe_crops_business_images(width: int, height: int) -> None:
    """横图、竖图和方图都必须进入业务图片槽并保留裁切。"""

    page = _renderer().render(
        template_id="template_19",
        semantic_slides=[{
            "type": "content",
            "data": {"title": "图文探针", "items": _items(1)},
            "images": _images(1, width=width, height=height),
        }],
        task_id=f"template-19-probe-image-{width}-{height}",
        fallback_title="图文探针",
    )["slides"][0]
    image = next(element for element in page["elements"] if element.get("imageType") == "content")
    assert image["src"].startswith("https://example.invalid/")
    assert image.get("clip", {}).get("shape") == "rect"
    assert (image.get("originalWidth"), image.get("originalHeight")) == (width, height)
    decorations = [element for element in page["elements"] if element.get("imageType") == "decoration"]
    assert decorations and all(element["src"].startswith("/api/data/template_19_asset_") for element in decorations)


def test_template_19_probe_rejects_missing_image_dimensions() -> None:
    with pytest.raises(TemplateRenderError):
        _renderer().render(
            template_id="template_19",
            semantic_slides=[{
                "type": "content",
                "data": {"title": "缺少尺寸", "items": _items(1)},
                "images": [{"src": "https://example.invalid/no-size.jpg"}],
            }],
            task_id="template-19-probe-missing-size",
            fallback_title="缺少尺寸",
        )


def test_template_19_probe_rejects_extra_images() -> None:
    with pytest.raises(TemplateRenderError, match="图片数量超过内容项数量"):
        _renderer().render(
            template_id="template_19",
            semantic_slides=[{
                "type": "content",
                "data": {"title": "图片数量", "items": _items(1)},
                "images": _images(2),
            }],
            task_id="template-19-probe-extra-images",
            fallback_title="图片数量",
        )


def test_template_19_production_inventory_matches_goal() -> None:
    """生产模板必须精确覆盖规划的18页和五种页面类型。"""

    template = _template()
    counts = {
        kind: sum(slide["type"] == kind for slide in template["slides"])
        for kind in ("cover", "contents", "transition", "content", "end")
    }
    assert template["metadata"]["buildStage"] == "production"
    assert set(template["metadata"]["mvpSlideIds"]) == MVP_IDS
    assert set(template["metadata"]["productionSlideIds"]) == PRODUCTION_IDS
    assert {slide["id"] for slide in template["slides"]} == PRODUCTION_IDS
    assert counts == {"cover": 2, "contents": 6, "transition": 2, "content": 6, "end": 2}


@pytest.mark.parametrize(
    ("stage", "expected"),
    [("probe", 3), ("sample", 5), ("mvp", 12), ("production", 18)],
)
def test_template_19_builder_is_deterministic(tmp_path: Path, stage: str, expected: int) -> None:
    outputs = [tmp_path / f"{stage}-{index}.json" for index in (1, 2)]
    for output in outputs:
        result = subprocess.run(
            ["node", str(BUILDER_PATH), "--stage", stage, str(output)],
            cwd=REPOSITORY_ROOT,
            capture_output=True,
            text=True,
            timeout=30,
            check=False,
        )
        assert result.returncode == 0, result.stderr
    first = json.loads(outputs[0].read_text(encoding="utf-8"))
    second = json.loads(outputs[1].read_text(encoding="utf-8"))
    assert first == second
    assert len(first["slides"]) == expected
    assert first["metadata"]["buildStage"] == stage


def test_template_19_checked_in_production_matches_builder(tmp_path: Path) -> None:
    """仓库发布 JSON 必须与确定性构建器的 production 输出逐字节一致。"""

    output = tmp_path / "template_19_production.json"
    result = subprocess.run(
        ["node", str(BUILDER_PATH), "--stage", "production", str(output)],
        cwd=REPOSITORY_ROOT,
        capture_output=True,
        text=True,
        timeout=30,
        check=False,
    )
    assert result.returncode == 0, result.stderr
    assert output.read_bytes() == TEMPLATE_PATH.read_bytes()


def test_template_19_roundtrip_verifier_rejects_unrelated_single_slide_deck(tmp_path: Path) -> None:
    """无关的一页 PPTX 不能再凭“非空”冒充生产模板往返验收通过。"""

    input_path = tmp_path / "unrelated-single-slide.pptx"
    output_path = tmp_path / "roundtrip-summary.json"
    deck = PowerPointPresentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(input_path)

    result = subprocess.run(
        ["node", str(ROUNDTRIP_VERIFIER_PATH), str(input_path), str(output_path)],
        cwd=REPOSITORY_ROOT,
        capture_output=True,
        text=True,
        timeout=30,
        check=False,
    )
    summary = json.loads(output_path.read_text(encoding="utf-8"))
    assert result.returncode == 1
    assert summary["status"] == "FAIL"
    assert summary["slideCount"] == 1
    assert "SLIDE_COUNT_MISMATCH" in summary["failures"]
    assert "EDITED_TITLE_MISSING" in summary["failures"]


def test_template_19_assets_are_exact_external_and_valid() -> None:
    """九项发布素材必须全部被引用并满足尺寸、模式、Alpha和体积契约。"""

    template = _template()
    referenced = {
        element["src"].rsplit("/", 1)[-1]
        for slide in template["slides"]
        for element in slide["elements"]
        if element.get("type") == "image"
        and element.get("imageType") == "decoration"
        and "template_19_asset_" in element.get("src", "")
    }
    published = {path.name for path in TEMPLATE_ROOT.glob("template_19_asset_*")}
    assert referenced == published == set(ASSET_CONTRACT)
    for filename, (size, mode, limit) in ASSET_CONTRACT.items():
        path = TEMPLATE_ROOT / filename
        with Image.open(path) as image:
            assert image.size == size
            assert image.mode == mode
            if mode == "RGBA":
                assert image.getchannel("A").getextrema() == (0, 255)
        assert path.stat().st_size <= limit


def test_template_19_ids_fonts_paths_and_rights_are_clean() -> None:
    template = _template()
    slide_ids = [slide["id"] for slide in template["slides"]]
    element_ids = [element["id"] for slide in template["slides"] for element in slide["elements"]]
    serialized = json.dumps(template, ensure_ascii=False)
    assert len(slide_ids) == len(set(slide_ids))
    assert len(element_ids) == len(set(element_ids))
    assert "data:image/" not in serialized and "file://" not in serialized
    assert "template_10_asset_" not in serialized and "template_18_asset_" not in serialized
    for forbidden in (
        "汉仪家书简",
        "迷你简卡通",
        "全字库正楷体",
        "站酷快乐体",
        "小笼包",
        "2018.10",
        "POWERPOINT",
    ):
        assert forbidden not in serialized
    assert "animations" not in serialized
    assert '"chart"' not in serialized.lower()
    assert all(slide.get("sourceReferenceSlides") for slide in template["slides"])


def test_template_19_preserves_reference_specific_visual_contract() -> None:
    template = _template()
    by_id = {slide["id"]: slide for slide in template["slides"]}
    cover = by_id["cover-botanical-frame"]
    assert any(element.get("src", "").endswith("template_19_asset_bg_cover_v1.jpg") for element in cover["elements"])
    contents = by_id["contents-4"]
    assert any(element.get("src", "").endswith("template_19_asset_eucalyptus_corner_v1.png") for element in contents["elements"])
    transition = by_id["transition-eucalyptus"]
    assert any(element.get("src", "").endswith("template_19_asset_bg_section_v1.jpg") for element in transition["elements"])
    assert any(_slot_type(element) == "partNumber" for element in transition["elements"])
    end = by_id["end-botanical-frame"]
    assert any(element.get("src", "").endswith("template_19_asset_bg_end_v1.jpg") for element in end["elements"])


@pytest.mark.parametrize("count", [2, 3, 4, 5, 6, 10])
def test_template_19_selects_exact_contents_capacity(count: int) -> None:
    values = [f"议题 {index}" for index in range(1, count + 1)]
    page = _renderer().render(
        template_id="template_19",
        semantic_slides=[{"type": "contents", "data": {"items": values}}],
        task_id=f"template-19-contents-{count}",
        fallback_title="目录",
    )["slides"][0]
    assert page["templateSlideId"] == f"contents-{count}"
    assert sum(_slot_type(element) == "item" for element in page["elements"]) == count


def test_template_19_paginates_eleven_contents_without_loss() -> None:
    values = [f"议题 {index}" for index in range(1, 12)]
    document = _renderer().render(
        template_id="template_19",
        semantic_slides=[{"type": "contents", "data": {"items": values}}],
        task_id="template-19-contents-eleven",
        fallback_title="目录",
    )
    rendered = [
        _plain_text(element)
        for slide in document["slides"]
        for element in slide["elements"]
        if _slot_type(element) == "item"
    ]
    assert len(document["slides"]) == 2
    assert rendered == values


@pytest.mark.parametrize(
    ("count", "expected"),
    [(1, "content-statement-1"), (2, "content-text-2"), (3, "content-text-3"), (4, "content-text-4")],
)
def test_template_19_selects_exact_text_capacity(count: int, expected: str) -> None:
    page = _renderer().render(
        template_id="template_19",
        semantic_slides=[{"type": "content", "data": {"title": "核心观点", "items": _items(count)}}],
        task_id=f"template-19-text-{count}",
        fallback_title="核心观点",
    )["slides"][0]
    assert page["templateSlideId"] == expected
    assert sum(_slot_type(element) == "item" for element in page["elements"]) == count
    assert not any(element.get("imageType") == "content" for element in page["elements"])


def test_template_19_metrics_selects_editable_metrics_layout() -> None:
    page = _renderer().render(
        template_id="template_19",
        semantic_slides=[{
            "type": "content",
            "data": {
                "title": "四项指标",
                "items": [
                    {"kind": kind, "title": f"指标 {index}", "text": f"{70 + index * 5}%"}
                    for index, kind in enumerate(("metric", "number", "stat", "metric"), 1)
                ],
            },
        }],
        task_id="template-19-metrics",
        fallback_title="四项指标",
    )["slides"][0]
    assert page["templateSlideId"] == "content-metrics-4"
    assert page.get("layoutKind") == "metrics"
    assert not any(element.get("type") == "chart" for element in page["elements"])


def test_template_19_paginates_eight_items_without_reordering() -> None:
    values = _items(8)
    document = _renderer().render(
        template_id="template_19",
        semantic_slides=[{"type": "content", "data": {"title": "八项内容", "items": values}}],
        task_id="template-19-eight-items",
        fallback_title="八项内容",
    )
    rendered = [
        _plain_text(element)
        for slide in document["slides"]
        for element in slide["elements"]
        if _slot_type(element) == "itemTitle"
    ]
    assert rendered == [f"核心要点{index:02d}" for index in range(1, 9)]
    source_titles = [
        item.get("sourceTitle")
        for slide in document["slides"]
        for item in slide.get("sourceData", {}).get("items", [])
    ]
    if source_titles:
        assert source_titles == [item["title"] for item in values]


def test_template_19_long_body_is_split_without_loss() -> None:
    body = "水彩绿植正文必须完整保留并保持顺序。" * 18
    document = _renderer().render(
        template_id="template_19",
        semantic_slides=[{
            "type": "content",
            "data": {"title": "长正文", "items": [{"title": "完整说明", "text": body}]},
        }],
        task_id="template-19-long-body",
        fallback_title="长正文",
    )
    rendered = "".join(
        _plain_text(element)
        for slide in document["slides"]
        for element in slide["elements"]
        if _slot_type(element) == "item"
    )
    assert len(document["slides"]) > 1
    assert rendered == body


def test_template_19_cover_and_transition_variants_are_reachable() -> None:
    without_image = _renderer().render(
        template_id="template_19",
        semantic_slides=[{"type": "cover", "data": {"title": "自然表达", "text": "清晰沟通"}}],
        task_id="template-19-cover-no-image",
        fallback_title="自然表达",
    )["slides"][0]
    with_image = _renderer().render(
        template_id="template_19",
        semantic_slides=[{
            "type": "cover",
            "data": {"title": "自然表达", "text": "清晰沟通"},
            "images": _images(1),
        }],
        task_id="template-19-cover-image",
        fallback_title="自然表达",
    )["slides"][0]
    transitions = _renderer().render(
        template_id="template_19",
        semantic_slides=[
            {"type": "transition", "data": {"title": "第一章节", "text": "说明一"}},
            {"type": "transition", "data": {"title": "第二章节", "text": "说明二"}},
        ],
        task_id="template-19-transitions",
        fallback_title="章节",
    )["slides"]
    assert without_image["templateSlideId"] == "cover-botanical-frame"
    assert with_image["templateSlideId"] == "cover-image"
    assert {slide["templateSlideId"] for slide in transitions} == {"transition-eucalyptus", "transition-fern"}


def test_template_19_cover_accepts_real_agent_fifty_character_summary() -> None:
    summary = ("水彩绿植封面副标题" + "需完整保持参考样式与信息层级" * 3)[:50]
    assert len(summary) == 50
    source = next(slide for slide in _template()["slides"] if slide["id"] == "cover-botanical-frame")
    panel = next(element for element in source["elements"] if element["id"] == "t19-cover-botanical-frame-title-panel-002")
    content_slot = next(element for element in source["elements"] if element["id"] == "t19-cover-botanical-frame-content-005")
    # 扩容后的副标题槽仍必须完整留在标题面板内，保持参考稿的封面框景关系。
    assert content_slot["top"] + content_slot["height"] <= panel["top"] + panel["height"]
    page = _renderer().render(
        template_id="template_19",
        semantic_slides=[{"type": "cover", "data": {"title": "自然表达", "text": summary}}],
        task_id="template-19-cover-fifty-character-summary",
        fallback_title="自然表达",
    )["slides"][0]
    rendered = [
        _plain_text(element)
        for element in page["elements"]
        if _slot_type(element) == "content"
    ]
    assert page["templateSlideId"] == "cover-botanical-frame"
    assert rendered == [summary]


def test_template_19_image_cover_accepts_real_agent_fifty_character_summary() -> None:
    summary = ("水彩绿植封面副标题" + "需完整保持参考样式与信息层级" * 3)[:50]
    page = _renderer().render(
        template_id="template_19",
        semantic_slides=[{
            "type": "cover",
            "data": {"title": "自然表达", "text": summary},
            "images": _images(1),
        }],
        task_id="template-19-image-cover-fifty-character-summary",
        fallback_title="自然表达",
    )["slides"][0]
    rendered = [
        _plain_text(element)
        for element in page["elements"]
        if _slot_type(element) == "content"
    ]
    assert page["templateSlideId"] == "cover-image"
    assert rendered == [summary]


@pytest.mark.parametrize("count", [0, 1, 2, 3])
def test_template_19_selects_end_layout_by_action_count(count: int) -> None:
    items = [f"行动 {index}" for index in range(1, count + 1)]
    page = _renderer().render(
        template_id="template_19",
        semantic_slides=[{"type": "end", "data": {"title": "下一步", "items": items}}],
        task_id=f"template-19-end-{count}",
        fallback_title="下一步",
    )["slides"][0]
    assert page["templateSlideId"] == ("end-botanical-frame" if count == 0 else "end-action")
    assert sum(_slot_type(element) == "item" for element in page["elements"]) == count


def test_template_19_rejects_more_than_three_end_actions() -> None:
    with pytest.raises(TemplateRenderError, match="不能超过 3 项"):
        _renderer().render(
            template_id="template_19",
            semantic_slides=[{
                "type": "end",
                "data": {"title": "下一步", "items": [f"行动 {index}" for index in range(4)]},
            }],
            task_id="template-19-end-too-many",
            fallback_title="下一步",
        )


def test_template_19_eighty_item_outline_stays_dense() -> None:
    semantic = []
    for page in range(1, 21):
        # 前十页使用短说明，后十页使用真实长度正文；固定 80 项应形成约 30 页而不退化为单项页。
        body = "真实业务说明。" if page <= 10 else "真实业务说明需要完整保留观点依据执行方法和后续行动。" * 2
        semantic.append({
            "type": "content",
            "data": {
                "title": f"内容页 {page}",
                "items": [
                    {"title": f"完整业务标题第{index}项需要安全展示", "text": body}
                    for index in range(1, 5)
                ],
            },
        })
    document = _renderer().render(
        template_id="template_19",
        semantic_slides=semantic,
        task_id="template-19-eighty-items",
        fallback_title="八十项容量",
    )
    counts = [
        sum(_slot_type(element) == "item" for element in slide["elements"])
        for slide in document["slides"]
        if slide.get("type") == "content"
    ]
    assert 28 <= len(document["slides"]) <= 35
    assert sum(count == 1 for count in counts) / max(1, len(counts)) < 0.2


def test_template_19_cover_and_registration_are_valid() -> None:
    cover = TEMPLATE_ROOT / "template_19.jpg"
    with Image.open(cover) as image:
        assert image.size == (960, 540)
        assert image.mode == "RGB"
        assert image.format == "JPEG"
    assert cover.stat().st_size < 150_000
    main_source = (TEMPLATE_ROOT.parent / "main.py").read_text(encoding="utf-8")
    assert main_source.count('{ "name": "水彩绿植轻商务", "id": "template_19"') == 1


def test_template_19_main_api_registration_and_cover_route() -> None:
    environment = {
        key: os.environ[key]
        for key in ("PATH", "PATHEXT", "SYSTEMROOT", "TEMP", "TMP", "WINDIR")
        if key in os.environ
    }
    environment.update({
        "APP_ENV": "test",
        "PERSISTENCE_ENABLED": "false",
        "SSO_ENABLED": "false",
        "BILLING_ENABLED": "false",
        "STORAGE_ENABLED": "false",
        "TASK_WORKER_ENABLED": "false",
        "RELEASE_CHANNEL": "test",
        "RELEASE_COMMIT": "template-19-test",
    })
    result = subprocess.run(
        [
            sys.executable,
            "-c",
            (
                "import json, dotenv; dotenv.load_dotenv=lambda *a, **k: False; import main; "
                "from fastapi.testclient import TestClient; client=TestClient(main.app); "
                "templates=client.get('/templates'); cover=client.get('/data/template_19.jpg'); "
                "items=templates.json()['data']; target=[item for item in items if item['id']=='template_19']; "
                "print(json.dumps({'templates_status':templates.status_code,'target':target,"
                "'unique':len({item['id'] for item in items})==len(items),"
                "'cover_status':cover.status_code,'cover_type':cover.headers.get('content-type'),"
                "'cover_bytes':len(cover.content)}))"
            ),
        ],
        cwd=REPOSITORY_ROOT / "backend/main_api",
        env=environment,
        capture_output=True,
        text=True,
        timeout=20,
        check=False,
    )
    assert result.returncode == 0, result.stderr
    assert json.loads(result.stdout.strip().splitlines()[-1]) == {
        "templates_status": 200,
        "target": [{
            "name": "水彩绿植轻商务",
            "id": "template_19",
            "cover": "/api/data/template_19.jpg",
        }],
        "unique": True,
        "cover_status": 200,
        "cover_type": "image/jpeg",
        "cover_bytes": (TEMPLATE_ROOT / "template_19.jpg").stat().st_size,
    }
